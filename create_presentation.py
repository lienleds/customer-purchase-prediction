from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

REPO_ROOT = Path(__file__).resolve().parent
OUTPUT_FILE = REPO_ROOT / "customer_purchase_prediction.pptx"
DATASET_FILE = REPO_ROOT / "Social_Network_Ads.csv"
REPO_NAME = "lienleds/customer-purchase-prediction"
AUTHOR = "@lienleds"

BLUE = RGBColor(25, 92, 167)
GREEN = RGBColor(38, 166, 154)
ORANGE = RGBColor(245, 124, 0)
DARK = RGBColor(33, 37, 41)
MUTED = RGBColor(92, 102, 112)
LIGHT = RGBColor(244, 248, 251)
WHITE = RGBColor(255, 255, 255)
BORDER = RGBColor(205, 214, 223)

TITLE_FONT = "Aptos Display"
BODY_FONT = "Aptos"


def dataset_size(default: int = 400) -> int:
    if not DATASET_FILE.exists():
        return default
    with DATASET_FILE.open(newline="", encoding="utf-8") as handle:
        rows = sum(1 for _ in csv.reader(handle)) - 1
    return rows if rows > 0 else default


def add_slide_title(slide, title: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.7))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = TITLE_FONT
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = DARK
    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.88), Inches(12.3), Inches(0.35))
        tf = subtitle_box.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = subtitle
        run.font.name = BODY_FONT
        run.font.size = Pt(11)
        run.font.color.rgb = MUTED


def add_banner(slide, color: RGBColor) -> None:
    banner = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
    banner.fill.solid()
    banner.fill.fore_color.rgb = color
    banner.line.fill.background()


def add_bullets(slide, left: float, top: float, width: float, height: float, bullets: Iterable[str], accent: RGBColor) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for index, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = BODY_FONT
        p.font.size = Pt(18)
        p.font.color.rgb = DARK
        p.bullet = True
        p.space_after = Pt(8)
    accent_bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left - 0.18), Inches(top), Inches(0.08), Inches(height))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()


def add_card(slide, left: float, top: float, width: float, height: float, title: str, body: list[str], accent: RGBColor) -> None:
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT
    card.line.color.rgb = BORDER

    title_box = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.15), Inches(width - 0.4), Inches(0.35))
    p = title_box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = BODY_FONT
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = accent

    body_box = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.55), Inches(width - 0.35), Inches(height - 0.7))
    tf = body_box.text_frame
    tf.word_wrap = True
    for index, item in enumerate(body):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = BODY_FONT
        p.font.size = Pt(15)
        p.font.color.rgb = DARK
        p.bullet = True
        p.space_after = Pt(6)


def add_flow_box(slide, left: float, top: float, width: float, height: float, label: str, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = WHITE
    text_frame = shape.text_frame
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.name = BODY_FONT
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = WHITE


def add_arrow(slide, start_x: float, start_y: float, end_x: float, end_y: float) -> None:
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(start_x),
        Inches(start_y),
        Inches(end_x),
        Inches(end_y),
    )
    connector.line.color.rgb = MUTED
    connector.line.width = Pt(2)
    connector.line.end_arrowhead = True


def build_presentation() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    size = dataset_size()

    # Slide 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_banner(slide, BLUE)
    hero = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(0.9), Inches(12.2), Inches(2.25))
    hero.fill.solid()
    hero.fill.fore_color.rgb = LIGHT
    hero.line.color.rgb = BORDER
    title = slide.shapes.add_textbox(Inches(0.9), Inches(1.25), Inches(11.3), Inches(0.9))
    p = title.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Customer Purchase Prediction\nUsing Logistic Regression"
    run.font.name = TITLE_FONT
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = DARK
    subtitle = slide.shapes.add_textbox(Inches(0.95), Inches(2.25), Inches(10.8), Inches(0.55))
    p = subtitle.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = f"Dataset: Social_Network_Ads.csv  |  Repository: {REPO_NAME}  |  Author: {AUTHOR}"
    run.font.name = BODY_FONT
    run.font.size = Pt(15)
    run.font.color.rgb = MUTED
    add_card(slide, 0.75, 3.55, 3.75, 2.5, "Training focus", ["Binary classification of Purchased (0/1)", "Core features: Age and EstimatedSalary"], BLUE)
    add_card(slide, 4.8, 3.55, 3.75, 2.5, "Why Logistic Regression", ["Simple, interpretable baseline", "Well suited to scaled tabular features"], GREEN)
    add_card(slide, 8.85, 3.55, 3.75, 2.5, "Deliverables", ["Python training script", "PowerPoint deck and reproducible generator"], ORANGE)

    # Slide 2
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_banner(slide, GREEN)
    add_slide_title(slide, "Requirement Understanding", "What the exercise asks us to deliver and verify")
    add_bullets(
        slide,
        0.95,
        1.45,
        5.45,
        4.75,
        [
            "Objective: predict whether a customer purchases a product.",
            "Dataset: Social_Network_Ads.csv for the training exercise.",
            "Input features: Age and EstimatedSalary.",
            "Target label: Purchased, where 0 = no purchase and 1 = purchase.",
            "Required evaluation: confusion matrix, accuracy, and feature interpretation.",
        ],
        GREEN,
    )
    add_card(slide, 7.1, 1.65, 5.1, 1.4, "Expected outputs", ["Model training result", "Evaluation summary", "Demo-ready presentation"], BLUE)
    add_card(slide, 7.1, 3.25, 5.1, 1.4, "Acceptance checks", ["Clear methodology", "Reproducible commands", "No unsupported metric claims"], ORANGE)
    add_card(slide, 7.1, 4.85, 5.1, 1.4, "Presentation emphasis", ["Concise bullets", "Readable visuals", "Professional training style"], GREEN)

    # Slide 3
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_banner(slide, BLUE)
    add_slide_title(slide, "Solution Overview", "End-to-end workflow used in the repository training script")
    add_bullets(
        slide,
        0.95,
        1.4,
        5.7,
        4.9,
        [
            "Explore the CSV, inspect data types, and confirm missing/duplicate status.",
            "Select Age and EstimatedSalary as predictors and Purchased as the target.",
            "Split the dataset into 80% training and 20% testing data.",
            "Normalize input features with StandardScaler.",
            "Train LogisticRegression on scaled training data.",
            "Evaluate predictions with accuracy, confusion matrix, and coefficient-based interpretation.",
        ],
        BLUE,
    )
    add_card(slide, 7.05, 1.6, 5.2, 1.6, "Repository implementation", ["Primary script: customer_purchase_prediction.py", "Generator script: create_presentation.py"], GREEN)
    add_card(slide, 7.05, 3.45, 5.2, 1.6, "Reusable pattern", ["Small tabular dataset", "Interpretable binary classifier baseline"], ORANGE)
    add_card(slide, 7.05, 5.3, 5.2, 0.95, "Key caution", ["Populate exact performance metrics from the script output if you want final production values."], BLUE)

    # Slide 4
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_banner(slide, ORANGE)
    add_slide_title(slide, "System Architecture", "Data flow from raw CSV to evaluation")
    flow_labels = [
        ("CSV\ninput", BLUE),
        ("Explore & clean", GREEN),
        ("Feature\nselection", ORANGE),
        ("Train/test\nsplit", BLUE),
        ("Scale with\nStandardScaler", GREEN),
        ("Logistic\nRegression", ORANGE),
        ("Prediction", BLUE),
        ("Evaluation", GREEN),
    ]
    x_positions = [0.45, 1.95, 3.55, 5.15, 6.9, 8.8, 10.45, 11.65]
    for (label, color), x in zip(flow_labels, x_positions):
        add_flow_box(slide, x, 2.9, 1.2, 1.0, label, color)
    for start, end in zip(x_positions, x_positions[1:]):
        add_arrow(slide, start + 1.2, 3.4, end, 3.4)
    add_card(slide, 1.0, 4.7, 4.0, 1.35, "Inputs", ["Age", "EstimatedSalary", "Purchased label for training"], BLUE)
    add_card(slide, 5.1, 4.7, 3.2, 1.35, "Outputs", ["Predicted class 0/1", "Probability estimates"], ORANGE)
    add_card(slide, 8.55, 4.7, 3.8, 1.35, "Verification", ["Accuracy score", "Confusion matrix review"], GREEN)

    # Slide 5
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_banner(slide, BLUE)
    add_slide_title(slide, "Accuracy, Benchmark, and Latency", "Use verified script output for exact runtime metrics")
    add_card(slide, 0.75, 1.45, 3.4, 2.1, "Benchmark facts", [f"Dataset size: {size} records", "Split: 80% train / 20% test", "Model: sklearn LogisticRegression"], BLUE)
    add_card(slide, 4.55, 1.45, 3.8, 2.1, "Exact metrics", ["Accuracy: <run script to fill>", "Training time: <measure during run>", "Prediction latency: <measure during run>"], ORANGE)
    add_card(slide, 8.75, 1.45, 3.8, 2.1, "How to obtain", ["Run: python create_presentation.py", "Run: python customer_purchase_prediction.py", "Copy verified results into final deck if needed"], GREEN)
    add_card(slide, 0.75, 4.0, 5.6, 1.8, "Confusion matrix definitions", ["TP: predicted purchase and actual purchase", "TN: predicted no purchase and actual no purchase", "FP: predicted purchase but actual no purchase", "FN: predicted no purchase but actual purchase"], BLUE)
    add_card(slide, 6.75, 4.0, 5.8, 1.8, "Reporting guidance", ["Do not guess exact scores.", "If execution is unavailable, keep placeholders and mention the command used to produce them."], ORANGE)

    # Slide 6
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_banner(slide, GREEN)
    add_slide_title(slide, "Conclusion and Demo Guidance", "What to emphasize during the training walkthrough")
    add_bullets(
        slide,
        0.95,
        1.4,
        5.9,
        4.9,
        [
            "Logistic Regression provides a clear baseline for customer purchase prediction.",
            "Age and EstimatedSalary influence the decision boundary after scaling.",
            "The workflow is reproducible and suitable for a short AI Foundation demo.",
            "Show the script run, generated analysis output, and this presentation in the video demo.",
            "Use the final slide to explain how confusion matrix results support the accuracy claim.",
        ],
        GREEN,
    )
    add_card(slide, 7.2, 1.75, 5.0, 1.5, "Commands to show", ["python -m pip install -r requirements.txt", "python create_presentation.py", "python customer_purchase_prediction.py"], BLUE)
    add_card(slide, 7.2, 3.65, 5.0, 1.3, "Files to highlight", ["customer_purchase_prediction.py", "create_presentation.py", "customer_purchase_prediction.pptx"], ORANGE)
    add_card(slide, 7.2, 5.35, 5.0, 0.9, "Demo tip", ["Open the PPTX after generation and replace placeholders only with verified script results."], GREEN)

    prs.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    output_path = build_presentation()
    print(f"Presentation created: {output_path}")
